"""Row-subset COPY of a BMR sheet — the "copy, do not construct" ruling.

Why this exists (register 2026-09-10, BMRtoAltID)
--------------------------------------------------
BMRtoAltID's input IS a BMR sheet, and the rows it wants to send on to
BMR-Review are rows of that sheet. Constructing them from parsed records
would lose every column BMRtoAltID never parses — and BMR-Review scores
some of those. A COPY cannot lose a column; a construction can. So the
reusable piece is not a writer but a subsetter, and it lives here rather
than in the consumer because the second consumer of a BMR sheet subset is
a matter of time, not of design.

Two answers BMRtoAltID asked for before this signature shipped (2026-09-11,
their P0), both baked in rather than documented:

* **There is no ``template=`` parameter.** Template family ("Non-Episodic"),
  file name and data-tab name ("Stand-Alone Works") are three strings for
  one thing, and the tab name is the one the caller already passes. The
  template is resolved from ``sheet_name`` through ``SHEET_TO_TEMPLATE``;
  a tab that is not a Template-22 data sheet is refused with the list.
* **``blank_columns=`` exists**, because the alternative is the consumer
  editing rows after the copy — the exact modification copying exists to
  prevent. Their sheets carry ``Candidates Found`` in ``Assigned EIDR ID``
  on 6,364 of 6,545 rows, written there by BMR-Review's own set-match. Is
  that a pre-fill? Read from BMR-Review's classifier (``classify_tamr``,
  2026-09-11): an ID is an assertion, ``NO MATCH`` is an assertion, and
  ANYTHING else — a sentinel or an empty cell alike — means "candidates:
  read them from ``Registration Errors & Notes``". So the sentinel alone is
  benign, but the notes beside it are not: a re-assessment that leaves the
  old notes in place evaluates the OLD candidate list instead of asking the
  fresh question, and a row that says ``NO MATCH`` asserts one. For a fresh
  assessment, blank BOTH ``Assigned EIDR ID`` and ``Registration Errors &
  Notes``. The blanking is explicit, by header name, refused for a name the
  sheet lacks (a typo must not be a silent no-op), and counted in the
  report.

Conformance is checked HERE, at emit time, against the shipped header row
(``Template.headers``), in two tiers. The operator's ruling (2026-09-11):
**a column that is optional in the registry, or that a child record can
inherit, may be empty or entirely missing from a BMR sheet without that
being an error** — Title Class and Alt ID Relation included. So:

* **Required**: the sheet mechanics every template carries (``Unique Row
  ID``, ``Operator's Notes``, ``Assigned EIDR ID``, ``Registration Errors &
  Notes``) plus, per template, the columns the registry REQUIRES and a
  record on that sheet can NEVER inherit. On a roots-only sheet
  (Stand-Alone Works, Compilations) that is the base structure — Structural
  Type, Mode, Referent Type, Title + language, Original Language 1, Release
  Date, Country of Origin 1, Publication Status, Approx Length. On a sheet
  that carries children it shrinks to the parent reference plus the
  creation-type block the schema makes mandatory: ``Edit Use`` / ``Color
  Type`` / ``In 3D`` and the Description, length and date an Edit must
  provide rather than inherit; ``Component Mode`` / ``Start Time`` /
  ``Content Duration`` for a Clip; ``Manif Class 1`` for a Manifestation;
  nothing but the parent for Episodics, since a Season or Episode inherits
  every base field (``eidr_core.inheritance.INHERITABLE_FIELDS``). Missing
  one of these is a ``TemplateMismatch``: the emit is refused and no output
  file is left behind.
* **Optional**: everything else the template ships — alternate titles,
  associated orgs, alt IDs and their relations, credits, classes, the
  IMDb / ISAN pair, every group-2+ member. Reported in
  ``SubsetReport.missing_optional``, never refused.

Measured before the ruling, with a stricter rule: of 294 Template-22 data
sheets under ``D:\\BMR``, "every shipped column present" would have refused
104 that the BMR tool accepted; the two-tier rule refuses exactly one,
which has no header row at all. The ruling only widens the optional tier,
so that result stands.

Extra columns and expanded families conform. "Is my input conforming
enough" is therefore answered per sheet, every time, rather than assumed
once — and answered against what the BMR tool actually accepts rather than
against the template's fullest form.

The copy is POSITIONAL and values-only: header rows 1..``DATA_START``-1 and
every kept row are written cell-for-cell into the same columns, in source
row order, onto a copy of the SOURCE workbook (so its other sheets,
validations and macros come along via the same transplant the writers
use). Nothing is renamed, reordered or re-mapped.
"""
from __future__ import annotations

import contextlib
import logging
import os
import shutil
import tempfile
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field

from eidr_core.bmr_io import DATA_START, fix_shared_strings, read_headers, transplant
from eidr_core.bmr_io.writer import SHEET_TO_TEMPLATE, TEMPLATES

log = logging.getLogger(__name__)

__all__ = ["TemplateMismatch", "SubsetReport", "SheetCheck", "FillReport", "subset_rows",
           "check_sheet", "fill_column", "required_headers"]

# Sheet mechanics: present on every template, written or read by the BMR
# tool itself, so a sheet without them is not a BMR sheet.
_MECHANICS = ("Unique Row ID", "Operator's Notes", "Assigned EIDR ID",
              "Registration Errors & Notes")

# Registry-required base structure of a ROOT record (common.xsd baseObjectData
# for a creation: StructuralType, Mode, ReferentType, ResourceName + lang,
# OriginalLanguage, ReleaseDate, CountryOfOrigin, Status, ApproximateLength
# all carry no minOccurs="0"). Every one of these is INHERITABLE by a child
# (eidr_core.inheritance.INHERITABLE_FIELDS), so they are required only on
# the two sheets that hold roots alone.
_ROOT_BASE = ("Structural Type", "Mode", "Referent Type", "Title", "Title Language",
              "Original Language 1", "Release Date", "Country of Origin 1",
              "Publication Status", "Approx Length")

# Per template: what the registry requires that a record on THAT sheet can
# never inherit. Registrant is required by the registry but supplied by the
# BMR tool's Configuration tab, so its column is not required here.
_REQUIRED: dict[str, tuple[str, ...]] = {
    "non_episodic": _MECHANICS + _ROOT_BASE,
    "compilation":  _MECHANICS + _ROOT_BASE,
    # Series rows are roots, but Season and Episode rows inherit every base
    # field, and the sheet-level check cannot tell the rows apart.
    "episodic":      _MECHANICS + ("Parent EIDR/Row ID",),
    # editInfoType: Parent, EditUse, ColorType, ThreeD have no minOccurs="0";
    # the schema annotation says an Edit must PROVIDE Description,
    # ApproximateLength and ReleaseDate rather than inherit them.
    "edit":          _MECHANICS + ("Parent EIDR/Row ID", "Edit Use", "Color Type", "In 3D",
                                   "Description", "Approx Length", "Release Date"),
    # clipInfoType: Parent, ComponentsMode, Start, Duration.
    "clip":          _MECHANICS + ("Parent EIDR/Row ID", "Component Mode", "Start Time",
                                   "Content Duration"),
    # manifestationInfoType: ManifestationClass maxOccurs=8, no minOccurs.
    "manifestation": _MECHANICS + ("Parent EIDR/Row ID", "Manif Class 1"),
}


def required_headers(template: str) -> tuple[str, ...]:
    """The shipped headers of ``template`` (a ``TEMPLATES`` key) whose
    absence makes a sheet NOT that template's data sheet: sheet mechanics
    plus the registry-required, never-inherited columns (see the module
    docstring for the ruling and the schema facts)."""
    shipped = TEMPLATES[template].headers          # KeyError on a non-key, by design
    req = _REQUIRED[template]
    missing = [h for h in req if h not in shipped]
    assert not missing, f"{template}: required columns not on the shipped template: {missing}"
    return tuple(h for h in shipped if h in req)


def _optional(header: str, template: str) -> bool:
    return header not in _REQUIRED[template]


class TemplateMismatch(ValueError):
    """The sheet is not (or is no longer) a conforming Template-22 data sheet."""


@dataclass
class SubsetReport:
    path: str
    template: str                 # TEMPLATES key resolved from the sheet name
    sheet: str
    rows_kept: int
    # Requested row numbers that do not exist on the source sheet. Not an
    # error: a caller subsetting from a stale finding list should see the
    # gap, not crash after writing nothing.
    rows_missing: list[int] = field(default_factory=list)
    # {column: non-empty values blanked} for each name in blank_columns.
    blanked: dict[str, int] = field(default_factory=dict)
    # Source headers that are not on the shipped template — preserved
    # verbatim (that is the point), listed so the caller knows they exist.
    extra_columns: list[str] = field(default_factory=list)
    # Shipped OPTIONAL columns the source lacks (companions, group 2+).
    # Reported, not refused — see the module docstring for the evidence.
    missing_optional: list[str] = field(default_factory=list)


@dataclass
class SheetCheck:
    """What ``check_sheet`` learned about a source sheet, without writing."""
    template: str                 # TEMPLATES key resolved from the sheet name
    sheet: str
    headers: dict[int, str]       # 1-based column -> header (header_map policy)
    extra_columns: list[str] = field(default_factory=list)
    missing_optional: list[str] = field(default_factory=list)


def _template_key(sheet_name: str) -> str:
    key = SHEET_TO_TEMPLATE.get(sheet_name)
    if key is None:
        raise TemplateMismatch(
            f"{sheet_name!r} is not a Template-22 data sheet; one of "
            f"{sorted(SHEET_TO_TEMPLATE)}")
    return key


def _inspect(wb, key: str, src_xlsx: str, sheet_name: str) -> SheetCheck:
    """The conformance rules, in ONE place: ``check_sheet`` runs them before
    a consumer starts an expensive run, and ``subset_rows`` runs the same
    function before it writes, so a sheet that passes the pre-flight cannot
    be refused at emit time for a reason the pre-flight did not know."""
    if sheet_name not in wb.sheetnames:
        raise TemplateMismatch(f"{src_xlsx} has no sheet {sheet_name!r} "
                               f"(sheets: {wb.sheetnames})")
    ws = wb[sheet_name]
    headers = read_headers(ws)
    present = set(headers.values())
    shipped = TEMPLATES[key].headers
    missing = [h for h in required_headers(key) if h not in present]
    if missing:
        raise TemplateMismatch(
            f"sheet {sheet_name!r} of {src_xlsx} is missing "
            f"{len(missing)} required Template-22 column(s): {missing}")
    return SheetCheck(
        template=key, sheet=sheet_name, headers=headers,
        extra_columns=[h for c, h in sorted(headers.items()) if h not in shipped],
        missing_optional=[h for h in shipped if h not in present and _optional(h, key)])


def check_sheet(src_xlsx: str, sheet_name: str) -> SheetCheck:
    """Pre-flight: would ``subset_rows`` accept ``sheet_name`` of ``src_xlsx``?

    Raises the same ``TemplateMismatch`` for the same reasons, and returns
    what the subsetter would have reported about the columns, without
    copying or writing anything. Added 0.33.0 for BMRtoAltID: its audit is a
    long mirror-bound run and the subset copy comes AFTER it, so a
    non-conforming source was refused an hour late. Asking first turns that
    into a two-second refusal. (A ``dry_run=`` flag on ``subset_rows`` was
    the other shape proposed; rejected because a dry run still needs a
    destination path and row list that mean nothing before the run.)
    """
    import openpyxl  # the `bmr` extra

    key = _template_key(sheet_name)
    if not os.path.exists(src_xlsx):
        raise FileNotFoundError(f"BMR sheet not found: {src_xlsx}")
    wb = openpyxl.load_workbook(src_xlsx, read_only=True, data_only=True)
    try:
        return _inspect(wb, key, src_xlsx, sheet_name)
    finally:
        wb.close()


@dataclass
class FillReport:
    path: str
    template: str
    sheet: str
    rows_filled: int
    # Keys of ``values`` that are not data rows on the sheet (below
    # DATA_START or beyond the last row). Reported, never dropped silently:
    # a caller's row numbers come from open_sheet on the same file, so a
    # miss here means something is wrong that they want to see.
    rows_absent: list[int] = field(default_factory=list)
    extra_columns: list[str] = field(default_factory=list)
    missing_optional: list[str] = field(default_factory=list)


def fill_column(src_xlsx: str, dst_xlsx: str, sheet_name: str, column: str,
                values: Mapping[int, object]) -> FillReport:
    """Copy ``src_xlsx`` to ``dst_xlsx`` and write ``values`` into ONE column
    of ``sheet_name``, keyed by absolute sheet row number.

    The sibling of ``subset_rows`` for the other half of a round trip
    (BMRtoAltID T5, 2026-09-12): a reviewed sheet comes back and the
    confirmed IDs go into ``Assigned EIDR ID`` so the existing modes can run
    on it. Same conformance rules as ``check_sheet`` (one ``_inspect``), same
    transplant so formatting and every other cell survive byte-for-byte,
    same refusal to leave a half-written file. Every data row is kept: this
    is not a subset. A ``column`` that is not on the header row is a
    ``ValueError`` (the caller named a column the sheet does not have -- a
    programming error, not a template mismatch).
    """
    import openpyxl  # the `bmr` extra

    key = _template_key(sheet_name)
    if not os.path.exists(src_xlsx):
        raise FileNotFoundError(f"BMR sheet not found: {src_xlsx}")

    shutil.copyfile(src_xlsx, dst_xlsx)
    try:
        wb = openpyxl.load_workbook(dst_xlsx, data_only=False)
        checked = _inspect(wb, key, src_xlsx, sheet_name)
        ws = wb[sheet_name]
        name_to_col = {h: c for c, h in checked.headers.items()}
        if column not in name_to_col:
            raise ValueError(f"column {column!r} is not on sheet {sheet_name!r}; "
                             f"headers: {sorted(name_to_col)}")
        col = name_to_col[column]
        report = FillReport(path=dst_xlsx, template=key, sheet=sheet_name, rows_filled=0,
                            extra_columns=checked.extra_columns,
                            missing_optional=checked.missing_optional)
        max_row = ws.max_row
        for r in sorted(int(k) for k in values):
            if r < DATA_START or r > max_row:
                report.rows_absent.append(r)
                continue
            ws.cell(r, col).value = values[r]
            report.rows_filled += 1

        out_dir = os.path.dirname(os.path.abspath(dst_xlsx)) or "."
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=out_dir)
        os.close(fd)
        try:
            wb.save(tmp)
            transplant(template_xlsx=dst_xlsx, edited_xlsx=tmp)
            fix_shared_strings(dst_xlsx)
        finally:
            with contextlib.suppress(OSError):
                if os.path.exists(tmp):
                    os.remove(tmp)
    except BaseException:
        with contextlib.suppress(OSError):
            os.remove(dst_xlsx)
        raise

    log.info("fill_column: %d row(s) of %r -> %s [%s]%s", report.rows_filled, column,
             dst_xlsx, sheet_name, f" absent={report.rows_absent}" if report.rows_absent else "")
    return report


def subset_rows(src_xlsx: str, dst_xlsx: str, sheet_name: str,
                keep_rows: Iterable[int], *,
                blank_columns: Sequence[str] = ()) -> SubsetReport:
    """Copy header rows and the ``keep_rows`` of ``sheet_name`` from
    ``src_xlsx`` into ``dst_xlsx``.

    ``keep_rows`` are ABSOLUTE sheet row numbers — what ``open_sheet`` /
    ``read_sheet`` emit and what a consumer's findings already carry. Rows
    below ``DATA_START`` are refused: the header rows are always kept and
    are not the caller's to choose.
    """
    import openpyxl  # the `bmr` extra

    key = _template_key(sheet_name)
    keep = sorted({int(r) for r in keep_rows})
    low = [r for r in keep if r < DATA_START]
    if low:
        raise ValueError(f"keep_rows must be data rows (>= {DATA_START}); header "
                         f"rows are always kept. Got {low}")
    if not os.path.exists(src_xlsx):
        raise FileNotFoundError(f"BMR sheet not found: {src_xlsx}")

    shutil.copyfile(src_xlsx, dst_xlsx)
    try:
        wb = openpyxl.load_workbook(dst_xlsx, data_only=False)
        checked = _inspect(wb, key, src_xlsx, sheet_name)
        ws = wb[sheet_name]
        headers = checked.headers
        name_to_col = {h: c for c, h in headers.items()}
        unknown = [b for b in blank_columns if b not in name_to_col]
        if unknown:
            raise ValueError(f"blank_columns not on sheet {sheet_name!r}: {unknown}")
        blank_cols = [name_to_col[b] for b in blank_columns]

        report = SubsetReport(
            path=dst_xlsx, template=key, sheet=sheet_name, rows_kept=0,
            blanked=dict.fromkeys(blank_columns, 0),
            extra_columns=checked.extra_columns,
            missing_optional=checked.missing_optional)

        max_row, max_col = ws.max_row, ws.max_column
        kept: list[list[object]] = []
        for r in keep:
            if r > max_row:
                report.rows_missing.append(r)
                continue
            values = [ws.cell(r, c).value for c in range(1, max_col + 1)]
            for name, c in zip(blank_columns, blank_cols, strict=True):
                v = values[c - 1]
                if v is not None and v != "":
                    report.blanked[name] += 1
                    values[c - 1] = None
            kept.append(values)

        # Clear the data area, then lay the kept rows down contiguously
        # from DATA_START in source order — positional, values only.
        if max_row >= DATA_START:
            for r in range(DATA_START, max_row + 1):
                for c in range(1, max_col + 1):
                    ws.cell(r, c).value = None
        for i, values in enumerate(kept, DATA_START):
            for c, v in enumerate(values, 1):
                if v is not None:
                    ws.cell(i, c).value = v
        report.rows_kept = len(kept)

        out_dir = os.path.dirname(os.path.abspath(dst_xlsx)) or "."
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=out_dir)
        os.close(fd)
        try:
            wb.save(tmp)
            transplant(template_xlsx=dst_xlsx, edited_xlsx=tmp)
            fix_shared_strings(dst_xlsx)
        finally:
            with contextlib.suppress(OSError):
                if os.path.exists(tmp):
                    os.remove(tmp)
    except BaseException:
        # A refused or failed emit leaves no half-written sheet for a
        # downstream tool to pick up as if it were complete.
        with contextlib.suppress(OSError):
            os.remove(dst_xlsx)
        raise

    log.info("subset_rows: %d of %d requested row(s) -> %s [%s]%s",
             report.rows_kept, len(keep), dst_xlsx, sheet_name,
             f" blanked={report.blanked}" if any(report.blanked.values()) else "")
    return report
