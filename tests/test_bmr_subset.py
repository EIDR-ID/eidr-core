"""Tests for ``subset_rows`` — copy, do not construct, and the two P0 answers.

Why (2026-09-11, BMRtoAltID P0.1 / P0.2)
-----------------------------------------
The subsetter's whole justification is that a copy cannot lose a column.
So the tests are about what SURVIVES: expanded families, out-of-template
columns, header rows, source order — read back from the workbook, never
from the return value. And about what is REFUSED before anything is
written: a non-conforming sheet, a tab that is not a Template-22 data
sheet, a blank column that is not there, a header row asked for as data.
"""
from __future__ import annotations

import os

import pytest

from eidr_core.bmr_io import DATA_START, HEADER_ROW, TEMPLATES
from eidr_core.bmr_io.subset import SubsetReport, TemplateMismatch, subset_rows

openpyxl = pytest.importorskip("openpyxl")

SHEET = "Stand-Alone Works"
CANON = list(TEMPLATES["non_episodic"].headers)
# A real-world shape: the shipped headers, plus an expanded family (a 4th
# Alt ID group), plus an out-of-template audit column after a blank
# separator — exactly what a BMR-Review assessed sheet looks like.
HEADERS = CANON[:CANON.index("IMDb")] + ["Alt ID 4", "Domain 4"] + CANON[CANON.index("IMDb"):]
EXTRA_COL = len(HEADERS) + 2          # one blank separator, then the extra
ROWS = 6


@pytest.fixture
def source(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.cell(1, 1).value = "EIDR Bulk Metadata Registration"
    ws.cell(2, 3).value = "banner cell in an odd column"
    for c, h in enumerate(HEADERS, 1):
        ws.cell(HEADER_ROW, c).value = h
    ws.cell(HEADER_ROW, EXTRA_COL).value = "Tamr EIDR ID"
    col = {h: c for c, h in enumerate(HEADERS, 1)}
    for i in range(ROWS):
        r = DATA_START + i
        ws.cell(r, col["Unique Row ID"]).value = f"R{i}"
        ws.cell(r, col["Title"]).value = f"Title {i}"
        ws.cell(r, col["Alt ID 4"]).value = f"alt4-{i}"
        ws.cell(r, col["Assigned EIDR ID"]).value = ("Candidates Found" if i % 2 else
                                                      ("NO MATCH" if i == 2 else None))
        ws.cell(r, col["Registration Errors & Notes"]).value = f"note {i}"
        ws.cell(r, EXTRA_COL).value = f"10.5240/T{i}"
    wb.create_sheet("Instructions").cell(1, 1).value = "keep me"
    path = tmp_path / "assessed.xlsx"
    wb.save(path)
    return str(path)


def _ws(path, sheet=SHEET):
    return openpyxl.load_workbook(path)[sheet]


def _col(name):
    return HEADERS.index(name) + 1


# ── what survives ────────────────────────────────────────────────────────

def test_kept_rows_land_contiguously_in_source_order_with_every_column(source, tmp_path):
    dst = str(tmp_path / "subset.xlsx")
    rep = subset_rows(source, dst, SHEET, keep_rows=[8, 4, 6])   # unordered on purpose
    ws = _ws(dst)
    assert [ws.cell(r, _col("Unique Row ID")).value for r in (4, 5, 6)] == ["R0", "R2", "R4"]
    assert ws.cell(7, _col("Unique Row ID")).value is None, "row 7 must be empty"
    # Expanded family and out-of-template column: preserved, same position.
    assert ws.cell(HEADER_ROW, _col("Alt ID 4")).value == "Alt ID 4"
    assert ws.cell(5, _col("Alt ID 4")).value == "alt4-2"
    assert ws.cell(HEADER_ROW, EXTRA_COL).value == "Tamr EIDR ID"
    assert ws.cell(6, EXTRA_COL).value == "10.5240/T4"
    # Sentinels copied VERBATIM when nothing was asked to be blanked.
    assert ws.cell(5, _col("Assigned EIDR ID")).value == "NO MATCH"
    assert rep == SubsetReport(path=dst, template="non_episodic", sheet=SHEET,
                               rows_kept=3, rows_missing=[],
                               blanked={}, extra_columns=["Alt ID 4", "Domain 4", "Tamr EIDR ID"])


def test_header_rows_and_other_sheets_come_along(source, tmp_path):
    dst = str(tmp_path / "subset.xlsx")
    subset_rows(source, dst, SHEET, keep_rows=[4])
    wb = openpyxl.load_workbook(dst)
    assert wb[SHEET].cell(1, 1).value == "EIDR Bulk Metadata Registration"
    assert wb[SHEET].cell(2, 3).value == "banner cell in an odd column"
    assert [wb[SHEET].cell(HEADER_ROW, c).value for c in range(1, len(HEADERS) + 1)] == HEADERS
    assert wb["Instructions"].cell(1, 1).value == "keep me"


def test_zero_rows_is_a_valid_headers_only_sheet(source, tmp_path):
    dst = str(tmp_path / "subset.xlsx")
    rep = subset_rows(source, dst, SHEET, keep_rows=[])
    assert rep.rows_kept == 0
    assert _ws(dst).cell(DATA_START, 1).value is None


def test_rows_beyond_the_sheet_are_reported_not_fatal(source, tmp_path):
    dst = str(tmp_path / "subset.xlsx")
    rep = subset_rows(source, dst, SHEET, keep_rows=[5, 500, 501])
    assert rep.rows_kept == 1 and rep.rows_missing == [500, 501]


# ── P0.1: blank_columns ──────────────────────────────────────────────────

def test_blank_columns_blanks_by_name_and_counts_only_what_was_there(source, tmp_path):
    dst = str(tmp_path / "subset.xlsx")
    rep = subset_rows(source, dst, SHEET, keep_rows=range(4, 10),
                      blank_columns=["Assigned EIDR ID", "Registration Errors & Notes"])
    ws = _ws(dst)
    for r in range(4, 10):
        assert ws.cell(r, _col("Assigned EIDR ID")).value is None
        assert ws.cell(r, _col("Registration Errors & Notes")).value is None
        assert ws.cell(r, _col("Title")).value is not None, "only the named columns"
    # Rows 0..5: odd -> Candidates Found (3), i==2 -> NO MATCH (1); notes on all 6.
    assert rep.blanked == {"Assigned EIDR ID": 4, "Registration Errors & Notes": 6}
    # The header cell itself is untouched.
    assert ws.cell(HEADER_ROW, _col("Assigned EIDR ID")).value == "Assigned EIDR ID"


def test_blank_column_that_is_not_there_is_refused_not_ignored(source, tmp_path):
    """'Assigned EIDR Id' must not be a silent no-op that ships 6,364 sentinels."""
    dst = str(tmp_path / "subset.xlsx")
    with pytest.raises(ValueError, match="blank_columns not on sheet"):
        subset_rows(source, dst, SHEET, keep_rows=[4], blank_columns=["Assigned EIDR Id"])
    assert not os.path.exists(dst)


# ── P0.2: the template is the tab, and conformance is checked at emit ────

def test_template_is_resolved_from_the_tab_name(source, tmp_path):
    rep = subset_rows(source, str(tmp_path / "s.xlsx"), SHEET, keep_rows=[4])
    assert rep.template == "non_episodic"


def test_a_tab_that_is_not_a_template_data_sheet_is_refused(source, tmp_path):
    dst = str(tmp_path / "subset.xlsx")
    with pytest.raises(TemplateMismatch, match="not a Template-22 data sheet"):
        subset_rows(source, dst, "Instructions", keep_rows=[4])
    with pytest.raises(TemplateMismatch, match="has no sheet 'Episodics'"):
        subset_rows(source, dst, "Episodics", keep_rows=[4])
    assert not os.path.exists(dst)


def test_a_sheet_missing_a_template_column_is_refused_and_leaves_no_file(source, tmp_path):
    ws_src = openpyxl.load_workbook(source)
    ws_src[SHEET].cell(HEADER_ROW, _col("Publication Status")).value = "Pub Status"
    ws_src.save(source)
    dst = str(tmp_path / "subset.xlsx")
    with pytest.raises(TemplateMismatch,
                       match=r"missing 1 required .*: \['Publication Status'\]"):
        subset_rows(source, dst, SHEET, keep_rows=[4])
    assert not os.path.exists(dst)


def test_missing_companion_and_group_two_columns_are_reported_not_refused(source, tmp_path):
    """Real member sheets omit these and the BMR tool takes them (2026-09-11
    survey of D:\\BMR: Candidates.xlsx has no Alt Title Class, the Movies
    baselines no Relation 2/3). Refusing them would refuse the round trip."""
    wb = openpyxl.load_workbook(source)
    ws = wb[SHEET]
    for name in ("Alt Title Class 1", "Alternate Title 2", "Alt Title Language 2",
                 "Alt Title Class 2", "Relation 3", "IMDb Relation"):
        ws.cell(HEADER_ROW, _col(name)).value = None
    wb.save(source)
    dst = str(tmp_path / "subset.xlsx")
    rep = subset_rows(source, dst, SHEET, keep_rows=[4, 5])
    assert rep.rows_kept == 2
    assert rep.missing_optional == ["Alt Title Class 1", "Alternate Title 2",
                                    "Alt Title Language 2", "Alt Title Class 2",
                                    "Relation 3", "IMDb Relation"]


def test_required_is_mechanics_plus_registry_required_never_inherited():
    """Operator ruling 2026-09-11: a column that is optional in the registry,
    or that a child can inherit, may be empty or missing without error."""
    from eidr_core.bmr_io.subset import required_headers
    req = required_headers("non_episodic")
    for s in ("Unique Row ID", "Assigned EIDR ID", "Registration Errors & Notes",
              "Structural Type", "Mode", "Referent Type", "Title", "Title Language",
              "Original Language 1", "Release Date", "Country of Origin 1",
              "Publication Status", "Approx Length"):
        assert s in req, s
    for q in ("Title Class", "Alternate Title 1", "Alt Title Class 1", "Associated Org 1",
              "Alt ID 1", "Domain 1", "Relation 1", "Director 1", "IMDb", "Registrant",
              "Language Mode 1", "Description"):
        assert q not in req, q
    # A sheet that holds children: every base field is inheritable.
    epi = required_headers("episodic")
    assert "Parent EIDR/Row ID" in epi and "Title" not in epi and "Release Date" not in epi
    # An Edit must PROVIDE its EditInfo, description, length and date.
    ed = required_headers("edit")
    for s in ("Edit Use", "Color Type", "In 3D", "Description", "Approx Length", "Release Date"):
        assert s in ed, s
    assert "Title" not in ed
    assert "Manif Class 1" in required_headers("manifestation")
    assert "Component Mode" in required_headers("clip")
    with pytest.raises(KeyError):
        required_headers("Stand-Alone Works")   # a key, not a tab name


def test_title_class_and_relation_may_be_missing_per_the_ruling(source, tmp_path):
    wb = openpyxl.load_workbook(source)
    ws = wb[SHEET]
    for name in ("Title Class", "Relation 1", "Relation 2", "Relation 3"):
        ws.cell(HEADER_ROW, _col(name)).value = None
    wb.save(source)
    rep = subset_rows(source, str(tmp_path / "s.xlsx"), SHEET, keep_rows=[4])
    assert rep.rows_kept == 1
    assert set(rep.missing_optional) == {"Title Class", "Relation 1", "Relation 2", "Relation 3"}


def test_every_template_declares_its_shipped_headers():
    for t in TEMPLATES.values():
        assert t.headers, t.key
        assert t.headers[0] == "Unique Row ID"
        assert t.headers[-3:] == ("Operator's Notes", "Assigned EIDR ID",
                                  "Registration Errors & Notes")
        assert len(set(t.headers)) == len(t.headers), f"{t.key}: duplicate header"


# ── 0.33.0: check_sheet is the subsetter's own rules, run before the copy ──

def test_check_sheet_reports_without_writing_and_agrees_with_subset_rows(source, tmp_path):
    from eidr_core.bmr_io import check_sheet
    before = sorted(os.listdir(tmp_path))
    chk = check_sheet(str(source), SHEET)
    assert sorted(os.listdir(tmp_path)) == before, "a pre-flight must not write"
    rep = subset_rows(str(source), str(tmp_path / "out.xlsx"), SHEET, [DATA_START])
    assert (chk.template, chk.sheet) == (rep.template, rep.sheet)
    assert chk.extra_columns == rep.extra_columns
    assert chk.missing_optional == rep.missing_optional
    assert chk.headers[_col("Unique Row ID")] == "Unique Row ID"


@pytest.mark.parametrize("breakage",
                         ["unknown-sheet-name", "no-such-sheet", "required-header-gone"])
def test_check_sheet_refuses_exactly_what_subset_rows_refuses(source, tmp_path, breakage):
    """The two must agree, or a source that passes the pre-flight is refused
    an hour later at emit time -- the failure check_sheet exists to prevent."""
    import openpyxl

    from eidr_core.bmr_io import check_sheet
    sheet = SHEET
    if breakage == "unknown-sheet-name":
        sheet = "Not A Template Sheet"
    elif breakage == "no-such-sheet":
        wb = openpyxl.load_workbook(source)
        wb[SHEET].title = "Renamed"
        wb.save(source)
    else:
        wb = openpyxl.load_workbook(source)
        wb[SHEET].cell(HEADER_ROW, _col("Unique Row ID")).value = None
        wb.save(source)
    with pytest.raises(TemplateMismatch) as pre:
        check_sheet(str(source), sheet)
    with pytest.raises(TemplateMismatch) as emit:
        subset_rows(str(source), str(tmp_path / "out.xlsx"), sheet, [DATA_START])
    assert str(pre.value) == str(emit.value)


# ── 0.35.1: the sheet's SHAPE is the subset's (BMRtoAltID, 2026-09-13) ────

def _sheet_xml(path, sheet=SHEET):
    """The worksheet part for ``sheet``: openpyxl writes parts in sheet order
    as xl/worksheets/sheetN.xml, and the transplant keeps that container."""
    import zipfile
    n = openpyxl.load_workbook(path, read_only=True).sheetnames.index(sheet) + 1
    with zipfile.ZipFile(path) as z:
        return z.read(f"xl/worksheets/sheet{n}.xml").decode("utf-8")


def test_the_output_sheet_has_exactly_the_kept_rows_and_a_matching_dimension(source, tmp_path):
    """A 5,000-row source subset to 856 rows still carried 5,003 <row>
    elements (4,144 empty) and declared A1:QY5003; a read-only reader saw
    max_row 5,003. Pinned: <row> count is header rows + rows_kept, and the
    <dimension> ends on that row."""
    import re
    out = tmp_path / "out.xlsx"
    rep = subset_rows(str(source), str(out), SHEET, [DATA_START])
    xml = _sheet_xml(out)
    rows = re.findall(r"<row ", xml)
    assert len(rows) == (DATA_START - 1) + rep.rows_kept, len(rows)
    dim = re.search(r'<dimension ref="([A-Z]+)(\d+):([A-Z]+)(\d+)"', xml)
    assert dim is not None, "no <dimension> element in the sheet XML"
    assert int(dim.group(4)) == (DATA_START - 1) + rep.rows_kept, dim.group(0)


def test_formula_text_survives_but_its_cached_value_does_not(source, tmp_path):
    """The documented behaviour of the openpyxl round-trip, pinned so a change
    is deliberate. A data_only reader sees None for a formula cell."""
    import openpyxl
    wb = openpyxl.load_workbook(source)
    ws = wb[SHEET]
    col = ws.max_column + 1
    ws.cell(HEADER_ROW, col).value = "Scratch"
    ws.cell(DATA_START, col).value = "=1+1"
    wb.save(source)
    out = tmp_path / "out.xlsx"
    subset_rows(str(source), str(out), SHEET, [DATA_START])
    assert openpyxl.load_workbook(out)[SHEET].cell(DATA_START, col).value == "=1+1"
    assert openpyxl.load_workbook(out, data_only=True)[SHEET].cell(DATA_START, col).value is None

