"""fill_column: one column changes, everything else is byte-for-byte the
source (BMRtoAltID T5, 2026-09-12)."""
from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import pytest

from eidr_core.bmr_io import DATA_START, HEADER_ROW, check_sheet, fill_column
from eidr_core.bmr_io.subset import TemplateMismatch
from tests.test_bmr_subset import SHEET, _col, source  # noqa: F401  (shared fixtures)


def _cells(path, sheet=SHEET):
    ws = openpyxl.load_workbook(path)[sheet]
    return {(r, c): ws.cell(r, c).value for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)}


def test_only_the_named_column_changes_and_the_input_is_untouched(source, tmp_path):  # noqa: F811 (shared fixture)
    src = Path(source)
    before = hashlib.sha256(src.read_bytes()).hexdigest()
    col = _col("Assigned EIDR ID")
    src_cells = _cells(source)
    rep = fill_column(str(source), str(tmp_path / "out.xlsx"), SHEET, "Assigned EIDR ID",
                      {DATA_START: "10.5240/AAAA-0000-0000-0000-0000-X",
                       DATA_START + 1: "10.5240/BBBB-0000-0000-0000-0000-Y"})
    assert hashlib.sha256(src.read_bytes()).hexdigest() == before, "input modified"
    assert rep.rows_filled == 2 and rep.rows_absent == []
    out = _cells(tmp_path / "out.xlsx")
    assert out[(DATA_START, col)] == "10.5240/AAAA-0000-0000-0000-0000-X"
    assert out[(DATA_START + 1, col)] == "10.5240/BBBB-0000-0000-0000-0000-Y"
    changed = {k for k in src_cells if src_cells[k] != out.get(k)}
    assert changed == {(DATA_START, col), (DATA_START + 1, col)}, changed
    assert out[(HEADER_ROW, col)] == "Assigned EIDR ID"


def test_rows_outside_the_data_area_are_reported_not_dropped(source, tmp_path):  # noqa: F811 (shared fixture)
    rep = fill_column(str(source), str(tmp_path / "out.xlsx"), SHEET, "Assigned EIDR ID",
                      {2: "header row", DATA_START: "ok", 10_000: "beyond the sheet"})
    assert rep.rows_filled == 1
    assert rep.rows_absent == [2, 10_000]


def test_an_unknown_column_is_a_value_error_and_leaves_no_file(source, tmp_path):  # noqa: F811 (shared fixture)
    with pytest.raises(ValueError, match="not on sheet"):
        fill_column(str(source), str(tmp_path / "out.xlsx"), SHEET, "No Such Column",
                    {DATA_START: 1})
    assert not (tmp_path / "out.xlsx").exists()


def test_conformance_is_check_sheets_own(source, tmp_path):  # noqa: F811 (shared fixture)
    wb = openpyxl.load_workbook(source)
    wb[SHEET].cell(HEADER_ROW, _col("Unique Row ID")).value = None
    wb.save(source)
    with pytest.raises(TemplateMismatch) as pre:
        check_sheet(str(source), SHEET)
    with pytest.raises(TemplateMismatch) as fill:
        fill_column(str(source), str(tmp_path / "out.xlsx"), SHEET, "Assigned EIDR ID", {})
    assert str(pre.value) == str(fill.value)
    assert not (tmp_path / "out.xlsx").exists()
